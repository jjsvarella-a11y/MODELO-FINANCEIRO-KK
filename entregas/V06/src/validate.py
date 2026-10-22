"""Structural checks, Python-vs-recalculation regression and financial invariants.
No claim of Microsoft Excel compatibility is made by this test runner.
"""
from pathlib import Path
import sys,json,math,re,posixpath
from zipfile import ZipFile
from lxml import etree as E
import openpyxl
OUT=Path(__file__).resolve().parents[1];REF=json.loads((OUT/'Referencia_Python_V06.json').read_text());NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def validate(path,include_legacy=False):
 results=[];failures=[];comparisons=0
 def check(name,ok,detail=None):
  results.append({'name':name,'passed':bool(ok),'detail':detail})
  if not ok:failures.append(name)
 with ZipFile(path) as z:
  check('ZIP CRC',z.testzip() is None);parts=set(z.namelist());issues=[];long=[];cellerrors=[];formula_count=0;dups=[];selected=[]
  styles=E.fromstring(z.read('xl/styles.xml'));nstyles=len(styles.find('m:cellXfs',NS));sst=E.fromstring(z.read('xl/sharedStrings.xml'));nstrings=len(sst)
  for p in sorted(parts):
   if not p.endswith(('.xml','.rels')):continue
   r=E.fromstring(z.read(p))
   if p.endswith('.rels'):
    owner='' if p=='_rels/.rels' else p.rsplit('/_rels/',1)[0]+'/'+p.rsplit('/_rels/',1)[1][:-5];ids=[]
    for rel in r:
     ids.append(rel.get('Id'))
     if rel.get('TargetMode')=='External':continue
     t=rel.get('Target','').split('#',1)[0];dest=posixpath.normpath(posixpath.join(posixpath.dirname(owner),t)).lstrip('/') if not t.startswith('/') else t.lstrip('/')
     if t and dest not in parts:issues.append([p,dest])
    if len(ids)!=len(set(ids)):issues.append([p,'duplicate Id'])
   if p.startswith('xl/worksheets/') and p.endswith('.xml'):
    if any(v.get('tabSelected') in ['1','true'] for v in r.findall('m:sheetViews/m:sheetView',NS)):selected.append(p)
    for row in r.findall('m:sheetData/m:row',NS):
     seen=set();last=0
     for c in row:
      addr=c.get('r');colnum=0
      for letter in re.match('[A-Z]+',addr).group():colnum=colnum*26+ord(letter)-64
      if addr in seen or colnum<=last:dups.append([p,addr])
      seen.add(addr);last=colnum
      if int(c.get('s','0'))>=nstyles:issues.append([p,addr,'bad style index'])
      if c.get('t')=='s' and int(c.findtext('m:v',namespaces=NS))>=nstrings:issues.append([p,addr,'bad string index'])
      if c.get('t')=='e':cellerrors.append([p,addr,c.findtext('m:v',namespaces=NS)])
      f=c.find('m:f',NS)
      if f is not None:
       formula_count+=1
       if len(f.text or '')>8192:long.append([p,addr,len(f.text or '')])
  check('Internal relationships, style/string indices',not issues,issues[:10]);check('Cell ordering and uniqueness',not dups,dups[:10]);check('Only one worksheet selected',len(selected)==1,selected);check('Excel formula length <=8192',not long,long[:10]);check('No cached formula errors',not cellerrors,cellerrors[:10])
  wbxml=E.fromstring(z.read('xl/workbook.xml'));sheets=wbxml.find('m:sheets',NS);check('Unique worksheet names',len({s.get('name') for s in sheets})==len(sheets));check('Nine dashboard names present',all(n in [s.get('name') for s in sheets] for n in REF['dashboards']));check('Homepage selected',sheets[0].get('name')=='EXEC V06');check('No external workbook links',not any(p.startswith('xl/externalLinks/') for p in parts));check('No VBA',not any('vbaProject' in p for p in parts))
 w=openpyxl.load_workbook(path,read_only=True,data_only=True);mismatches=[]
 cols=REF['main_columns']
 for idx,(label,expected) in enumerate(REF['scenarios'].items()):
  sn='V06 Motor' if idx==0 else f'V06 S{idx:02}'
  for m,(row,e) in enumerate(zip(w[sn].iter_rows(min_row=5,max_row=184,values_only=True),expected),1):
   for key,c in cols.items():
    actual=row[c];v=e[key];comparisons+=1
    if not isinstance(actual,(int,float)) or not math.isfinite(actual) or abs(actual-v)>max(.01,abs(v)*1e-9):mismatches.append([label,m,key,actual,v])
 check('All 25 x 180-month engines match Python',not mismatches,{'comparisons':comparisons,'mismatches':mismatches[:20]})
 # Independently recompute accounting identities, non-revolving usage and limits.
 inv=[]
 for label,rows in REF['scenarios'].items():
  totaldraw=0;lastcash=0;lastdebt=0
  for d in rows:
   totaldraw+=d['draw'];m=d['month']
   if abs(d['cash']-(lastcash+d['equity']+d['project_cf']+d['draw']-d['commitment_fee']-d['bank_fees']-d['interest']-d['principal']-d['sweep']-d['dividend']))>.01:inv.append([label,m,'cash'])
   if abs(d['debt']-(lastdebt+d['draw']-d['principal']-d['sweep']))>.01:inv.append([label,m,'debt'])
   if totaldraw>30000000.01:inv.append([label,m,'facility'])
   if d['dividend']>.01 and (m not in range(54,175,12) or d['cash']<d['reserve']-.01 or d['dividend']>d['legal_headroom']+.01 or d['dividend']>d['coverage_headroom']+.01):inv.append([label,m,'dividend constraint'])
   if d['owner_distribution']>max(0,d['res_cash'])+.01:inv.append([label,m,'owner cash'])
   lastcash=d['cash'];lastdebt=d['debt']
 check('Financial identities and constraints across scenarios',not inv,inv[:20])
 sys.path.insert(0,str(OUT/'src/v04/src'));import engine as legacy
 units=list(w['V06 Unidades'].iter_rows(min_row=5,max_row=98,values_only=True));check('94 unique apartment IDs and preserved source prices',len(units)==94 and len({str(x[0]) for x in units})==94 and all(str(x[0])==str(u['id']) and abs(x[7]-u['price'])<.01 for x,u in zip(units,legacy.UNITS)))
 if include_legacy:
  legacy_receipts=legacy.receipts(legacy.config())['units'];legacy_fail=[]
  for i,(row,expected) in enumerate(zip(w['Recebimentos'].iter_rows(min_row=5,max_row=98,max_col=181,values_only=True),legacy_receipts)):
   for month,(actual,value) in enumerate(zip(row[1:],expected),1):
    if abs((actual or 0)-value)>.01:legacy_fail.append([i,month,actual,value])
  check('16920 preserved V04 unit receipt values reconcile',not legacy_fail,{'comparisons':16920,'failures':legacy_fail[:10]})
 check('Construction delays do not extend bank draw window',all(d['draw']<.005 for d in REF['scenarios']['Construction +6M'] if d['month']>30))
 check('Operating tax exemptions do not erase development profit',all(REF['scenarios'][label][35]['annual_tax']>0 for label in ['IRPC 3 years','IRPC 5 years','IRPC full']))
 main=REF['main_rows'];check('94 apartment receipts total GDV',abs(sum(d['receipts'] for d in main)-50656336.37)<.01);check('CAPEX base 63168000',abs(sum(d['capex'] for d in main)-63168000)<.01);check('VAT base 8629200',abs(sum(d['vat_incurred'] for d in main)-8629200)<.01);check('Construction insurance 197400',abs(sum(d['construction_insurance'] for d in main)-197400)<.01);check('VAT six-month exposure benchmark',abs(max(d['vat_credit'] for d in main)-3077345.304)<.01);check('HOTUSA protected with no payment',all(d['hotusa_reserve']==1000000 and d['hotusa_payment']==0 for d in main));check('Combined stress exposes deficit',REF['metrics']['Combined stress']['gap']>0)
 # Verify fifteen blocks reconcile exactly to monthly distributions and aggregate investor cash flow.
 flows=list(w['V06 Fluxos Invest'].iter_rows(min_row=5,max_row=185,values_only=True));blocksum=[]
 for m,row in enumerate(flows):
  if abs(sum(row[20:35])-main[m-1]['dividend'])>.01 if m else abs(sum(row[20:35]))>.01:blocksum.append([m,'dividends'])
  if abs(sum(row[8:16])-row[16])>.01:blocksum.append([m,'new investor flow'])
 check('Fifteen-block cash-flow allocation',not blocksum,blocksum[:10])
 owners=list(w['V06 Blocos'].iter_rows(min_row=5,max_row=19,values_only=True));check('15 ownership blocks total 100%',len(owners)==15 and abs(sum(r[3] for r in owners)-1)<1e-12)
 # Summary IRR/NPV/MOIC must equal independent Python expected calculations.
 metric_errors=[]
 for i,(label,metrics) in enumerate(REF['metrics'].items(),5):
  row=next(w['V06 Cenarios'].iter_rows(min_row=i,max_row=i,values_only=True))
  for key,c in [('irr',10),('moic',11),('npv',12),('exit',9),('pnpv',19),('pirr',20)]:
   v=metrics[key];a=row[c]
   if isinstance(v,(int,float)):
    if not isinstance(a,(int,float)) or abs(a-v)>max(.01 if key in ['npv','exit','pnpv'] else 1e-7,abs(v)*1e-8):metric_errors.append([label,key,a,v])
   elif a!=v:metric_errors.append([label,key,a,v])
 check('Scenario return metrics match Python',not metric_errors,metric_errors[:10]);w.close()
 report={'tested_file':str(path),'passed':sum(x['passed'] for x in results),'failed':len(failures),'checks':results,'formula_count':formula_count,'sheet_count':len(sheets),'defined_name_count':len(wbxml.findall('m:definedNames/m:definedName',NS)),'chart_count':sum(bool(re.fullmatch(r'xl/charts/chart[0-9]+\.xml',p)) for p in parts),'style_count':nstyles,'workbook_bytes':Path(path).stat().st_size,'external_workbook_links':sum(p.startswith('xl/externalLinks/') and p.endswith('.xml') for p in parts),'native_excel_365':'NOT TESTED','recalculation_engine':'LibreOffice (supplementary computation, not Microsoft certification)'}
 target=OUT/('Validation_Final.json' if include_legacy else 'Validation_Modules.json');target.write_text(json.dumps(report,indent=2,ensure_ascii=False));print(json.dumps({'passed':report['passed'],'failed':report['failed'],'failures':failures,'file':str(target)},ensure_ascii=False));return report
if __name__=='__main__':
 r=validate(Path(sys.argv[1]),'--final' in sys.argv);raise SystemExit(r['failed']>0)
