from engine import *
from selected_policy import selected
import openpyxl,sys
path=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'validacao/V04_Referencia/out/V04_Referencia.xlsx';w=openpyxl.load_workbook(path,data_only=True,read_only=True);meta=json.loads((ROOT/'src/v04_risk_metadata.json').read_text());expected=json.loads((ROOT/'src/v04_risk_results.json').read_text());M=json.loads((ROOT/'src/build_metadata.json').read_text());out=[]
for item,r in zip(meta,expected):
 failures=[];count=0
 for row,e in zip(w[item['flow']].iter_rows(min_row=5,max_row=184,values_only=True),r['flow']):
  for key,col in M['flow_columns'].items():
   if isinstance(e[key],(int,float)):
    count+=1
    if not isinstance(row[col],(int,float)) or abs(row[col]-e[key])>max(.01,abs(e[key])*1e-9):failures.append(f'm{e["m"]} {key}: {row[col]} != {e[key]}')
 pol=selected(r)
 for row,e in zip(w[item['waterfall']].iter_rows(min_row=5,max_row=184,values_only=True),pol['rows']):
  expected_cells={6:e['dividends'][0],7:e['dividends'][1],8:e['dividends'][2],9:e['exit'][0],10:e['exit'][1],11:e['exit'][2],19:e['new_received'],20:e['new_outlay'],21:e['new_cf'],22:e['jscv_received'],23:e['jals_received'],24:e['promoter_cf']}
  for col,val in expected_cells.items():
   count+=1
   if not isinstance(row[col],(int,float)) or abs(row[col]-val)>max(.01,abs(val)*1e-9):failures.append(f'WFm{e["m"]} c{col}: {row[col]} != {val}')
 out.append({'case':item['code'],'ok':not failures,'comparisons':count,'failures':failures[:16]});print(json.dumps(out[-1],ensure_ascii=False),flush=True)
w.close();(ROOT/'entrega/Testes_Riscos_Integrados.json').write_text(json.dumps(out,ensure_ascii=False,indent=2));raise SystemExit(not all(x['ok'] for x in out))
