"""Executive PDF snapshot of final Excel caches, with real PDF navigation links."""
from engine import ROOT
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.utils import ImageReader
import openpyxl,math
from report_base_v04 import n,e,pc
P=ROOT/'entrega';w=openpyxl.load_workbook(P/'KK01_Modelo_Integrado_V04.xlsx',read_only=True,data_only=True);s=w['EXEC KK01'];a={c.coordinate:c.value for row in s.iter_rows(max_row=60,max_col=16) for c in row if hasattr(c,'coordinate')};risk=list(w['V04 Riscos'].iter_rows(min_row=5,max_row=18,values_only=True));ret={c.coordinate:c.value for row in w['V04 Retornos'].iter_rows(max_row=60,max_col=11) for c in row if hasattr(c,'coordinate')}
W,H=842,1191;c=canvas.Canvas(str(P/'KK01_Executive_Control_Center_V04.pdf'),pagesize=(W,H));c.setTitle('KK-01 | Executive Control Center V04');navy=colors.HexColor('#102E42');teal=colors.HexColor('#007F87');light=colors.HexColor('#EDF3F6');ink=colors.HexColor('#17364A');gold=colors.HexColor('#B78427')
def txt(x,y,t,size=10,color=ink,font='Helvetica'):c.setFillColor(color);c.setFont(font,size);c.drawString(x,y,str(t))
def wrap(x,y,t,width=114,size=9,color=ink):
 import textwrap
 for line in textwrap.wrap(t,width):txt(x,y,line,size,color);y-=size*1.35
 return y
c.setFillColor(navy);c.rect(0,H-122,W,122,fill=1,stroke=0);txt(32,H-46,'KK-01',29,colors.white,'Helvetica-Bold');txt(156,H-44,'EXECUTIVE CONTROL CENTER',20,colors.white,'Helvetica-Bold');txt(33,H-76,'V04 | Controlo estratégico, financeiro e operacional',13,colors.HexColor('#A4DCE0'));txt(33,H-100,'18/10/2026 | Referência projetada: Hotel Favorável · Banco 6% / 42M / 10A · Linha 30 M€',10,colors.white)
labels=[('Capital KK previsto','A7',e),('Quota paga equivalente','E7',pc),('Pico da dívida','I7',e),('Utilização acumulada','M7',e),('Caixa mínima','A12',e),('Dívida na abertura','E12',e),('Juros 180 meses','I12',e),('DSCR ano parcial','M12',lambda x:n(x)+' x'),('EBITDA ano completo','A17',e),('TIR novos efetiva 15A','E17',pc),('MOIC novos efetivo','I17',lambda x:n(x)+' x'),('VPL novos a 12%','M17',e)]
gap=10;cw=(W-64-3*gap)/4;ch=77
for i,(label,addr,fmt) in enumerate(labels):
 x=32+(i%4)*(cw+gap);y=H-149-(i//4)*(ch+10)-ch;c.setFillColor(light);c.roundRect(x,y,cw,ch,6,fill=1,stroke=0);txt(x+12,y+55,label,10);val=a[addr];color=colors.HexColor('#AD4D37') if isinstance(val,(int,float)) and val<0 else teal;txt(x+12,y+27,fmt(val),18,color,'Helvetica-Bold')
y=H-430;txt(32,y,'NAVEGAÇÃO | clique para abrir o relatório correspondente',11,ink,'Helvetica-Bold')
links=[('Investimento / CAPEX','Construcao_EAC'),('Real Estate / Comercial','Comercializacao_94_Apartamentos'),('Financiamento / Liquidez','Financiamento_Bancabilidade'),('Hotel + MICE','Hotel_Residential_Operator'),('Operador Internacional','Hotel_Residential_Operator'),('Exploração residencial','Hotel_Residential_Operator'),('Investidores / Acionistas','Estrutura_Acionista_Waterfall'),('Riscos / Auditoria','Auditoria')]
base='https://github.com/jjsvarella-a11y/SIM-JSCV-KYENU/blob/entrega-k01/entregas/V04/'
for i,(label,part) in enumerate(links):
 x=32+(i%4)*(cw+gap);yy=y-14-(i//4)*74-64;c.setFillColor(navy);c.roundRect(x,yy,cw,64,6,fill=1,stroke=0);c.drawImage(str(ROOT/'src/icons'/f'{chr(65+i)}.png'),x+9,yy+19,26,26,mask='auto');wrap(x+42,yy+40,label,width=22,size=9,color=colors.white);c.linkURL(base+'KK01_'+part+'_V04.pdf',(x,yy,x+cw,yy+64),relative=0)
y-=183
short=[('Workflow','Strategic_Master_Workflow'),('Pipeline','Equity_Investor_Pipeline'),('Completion M24','Financiamento_Bancabilidade'),('Decisões Board','Decisoes_Pendentes'),('Dicionário KPI','Dicionario_KPI'),('Riscos','Cenarios_Riscos')]
for i,(label,part) in enumerate(short):
 x=32+i*130;c.setFillColor(light);c.roundRect(x,y-28,120,29,4,fill=1,stroke=0);txt(x+9,y-17,label,10);c.linkURL(base+'KK01_'+part+'_V04.pdf',(x,y-28,x+120,y+1),relative=0)
y-=64;txt(32,y,'REMUNERAÇÃO E SEGURANÇA | hipóteses separadas de execução comprovada',11,ink,'Helvetica-Bold')
for i,(title,val) in enumerate([('Platinum convencional',pc(ret['D5'])),('Gold convencional',pc(ret['D6'])),('Silver convencional',pc(ret['D7'])),('Completion assegurada','Indisponível'),('EAC com Actual Cost','Indisponível'),('Prontidão documental','Indisponível')]):
 x=32+(i%3)*262;yy=y-15-(i//3)*65-55;c.setFillColor(light);c.roundRect(x,yy,252,55,4,fill=1,stroke=0);txt(x+12,yy+35,title,10);txt(x+12,yy+14,val,15,teal if i<3 else gold,'Helvetica-Bold')
y-=168;txt(32,y,'LEITURA PARA O BOARD',11,ink,'Helvetica-Bold');y-=22
for text in [
 'Estrutura reconciliada: 20 blocos, 100.000 ações. JSCV 1P + 2G = 40%; JALS 5S = 10%; novos 1P + 1G + 10S = 50%.',
 'As metas 10% / 6% / 5,5% são alcançáveis na proposta anual de waterfall. Direitos preferenciais dependem de aprovação.',
 'Completion projetada M24: '+n(risk[0][7])+' x. Recursos firmes e condições bancárias ainda não documentados.',
 'Investidores +12M e vendas +6M: insuficiência de reservas de '+e(risk[5][2])+'. A linha 30 M€ é back-up estratégico.',
 'VPL novos negativo à taxa de 12%. Retornos dependem da saída estimada; o modelo não constitui garantia de liquidez.',
 'Dados de execução, compromissos, pagamentos, licenças e HMA não fornecidos permanecem indisponíveis.'
]:
 y=wrap(40,y,'• '+text,width=124,size=10);y-=10
c.setStrokeColor(colors.HexColor('#CAD7DE'));c.line(32,52,W-32,52);txt(32,35,'Snapshot do Excel V04 recalculado | PDF estático; o Excel contém seletores, gráficos e notas auditáveis.',9);txt(W-47,35,'1',9);c.showPage();c.save();w.close();print('Executive PDF gerado,1 página A3,com links aos relatórios.')
