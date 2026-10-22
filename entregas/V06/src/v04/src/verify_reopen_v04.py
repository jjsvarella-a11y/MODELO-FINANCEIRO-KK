from engine import ROOT
import openpyxl,json,zipfile,xml.etree.ElementTree as E,re
P=ROOT/'entrega/KK01_Modelo_Integrado_V04.xlsx';Q=ROOT/'validacao/Reabertura_Final/KK01_Modelo_Integrado_V04.xlsx';a=openpyxl.load_workbook(P,read_only=True,data_only=True);b=openpyxl.load_workbook(Q,read_only=True,data_only=True);fail=[];count=0
for sn,maxrow,maxcol in [('Fluxo 180M',184,55),('V04 Retornos',60,11),('V04 Completion',23,3),('KPI',65,3),('Matriz Societaria V03',18,9)]:
 for ra,rb in zip(a[sn].iter_rows(max_row=maxrow,max_col=maxcol,values_only=True),b[sn].iter_rows(max_row=maxrow,max_col=maxcol,values_only=True)):
  for va,vb in zip(ra,rb):
   if isinstance(va,(int,float)):
    count+=1
    if not isinstance(vb,(int,float)) or abs(va-vb)>max(1e-8,abs(va)*1e-10):fail.append((sn,va,vb))
a.close();b.close();N='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}';errs=[];notes=0
with zipfile.ZipFile(Q) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.startswith('xl/worksheets/') and name.endswith('.xml') and re.search(rb'<c\b[^>]*\bt="e"',z.read(name)):errs.append(name)
  if name.startswith('xl/comments') and name.endswith('.xml'):
   for c in E.fromstring(z.read(name)).iter(N+'comment'):
    t=''.join(x.text or '' for x in c.iter(N+'t'))
    if 'Dicionário: V04 Dicionario KPI / ' in t:
     notes+=1
     if not all(x in t for x in ['Definição:','Interpretação:','AUDITORIA','Método:']):fail.append(('note',c.attrib['ref']))
report=dict(ok=not fail and not errs and notes==120,engine='LibreOfficeDev26.8',excel_desktop_tested=False,numeric_comparisons=count,kpi_notes_preserved=notes,formula_error_sheets=errs,failures=fail[:10],input_bytes=P.stat().st_size,reopened_bytes=Q.stat().st_size);(ROOT/'entrega/Verificacao_Reabertura_V04.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False,indent=2));raise SystemExit(not report['ok'])
