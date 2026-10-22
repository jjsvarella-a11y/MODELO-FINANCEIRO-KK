from pathlib import Path
import openpyxl,xlsxwriter
ROOT=Path(__file__).resolve().parents[1]
def extract():
 src=openpyxl.load_workbook('/workspace/kk01_efcc/final_recalc/out/V06_modulos.xlsx',read_only=True,data_only=True);ws=src['V06 Dicionario'];wb=xlsxwriter.Workbook(str(ROOT/'Dicionario_KPI.xlsx'));s=wb.add_worksheet('KPI PT EN');fmt=wb.add_format({'text_wrap':True,'valign':'top'});head=wb.add_format({'bold':True,'bg_color':'B31F2B','font_color':'FFFFFF','text_wrap':True})
 for r,row in enumerate(ws.iter_rows(min_row=4,values_only=True)):
  s.write_row(r,0,row,head if r==0 else fmt);s.set_row(r,42 if r else 28)
 s.set_column('A:C',26);s.set_column('D:M',44);s.freeze_panes(1,1);s.autofilter(0,0,r,12);wb.close();src.close();print('Dictionary snapshot exported')
if __name__=='__main__':extract()
